"""Ekstrakcja tekstu z DOCX (akapity + tabele)."""

from __future__ import annotations

import io

from docx import Document as DocxDocument

from backend.core.exceptions import ExtractionError

from .base import clean


def extract(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Nie udało się odczytać pliku DOCX ({exc}).") from exc
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = clean("\n".join(parts))
    if not text:
        raise ExtractionError("Plik DOCX jest pusty.")
    return text
