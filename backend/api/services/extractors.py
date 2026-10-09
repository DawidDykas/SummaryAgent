import io
import re

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text

from backend.core.exceptions import ExtractionError

__all__ = ["ACCEPT", "ExtractionError", "FORMATS", "extract_text", "format_of"]


FORMATS: dict[str, str] = {
    ".pdf": "PDF",
    ".docx": "DOCX",
    ".txt": "TXT",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".html": "HTML",
    ".htm": "HTML",
    ".rtf": "RTF",
}

ACCEPT = ",".join(f"*{ext}" for ext in FORMATS)


def format_of(file_name: str) -> str | None:
    dot = file_name.rfind(".")
    if dot == -1:
        return None
    return FORMATS.get(file_name[dot:].lower())


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Nie udało się odczytać pliku PDF ({exc}).") from exc
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = _clean("\n\n".join(pages))
    if not text:
        raise ExtractionError(
            "PDF nie zawiera warstwy tekstowej (prawdopodobnie skan) - nie da się go streszczyć."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Nie udało się odczytać pliku DOCX ({exc}).") from exc
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _clean("\n".join(parts))
    if not text:
        raise ExtractionError("Plik DOCX jest pusty.")
    return text


def _extract_html(data: bytes) -> str:
    soup = BeautifulSoup(_decode(data), "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    text = soup.get_text(separator="\n")
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return _clean(text)


def _extract_rtf(data: bytes) -> str:
    raw = data.decode("latin-1", errors="replace")
    try:
        text = rtf_to_text(raw, errors="ignore")
    except Exception as exc:
        raise ExtractionError(f"Nie udało się odczytać pliku RTF ({exc}).") from exc
    text = _clean(text)
    if not text:
        raise ExtractionError("Plik RTF jest pusty.")
    return text


def _extract_plain(data: bytes) -> str:
    text = _clean(_decode(data))
    if not text:
        raise ExtractionError("Plik jest pusty.")
    return text


def extract_text(file_name: str, data: bytes) -> tuple[str, str]:
    """Zwraca (format, tekst). Gdy format nieobsługiwany - ExtractionError."""
    fmt = format_of(file_name)
    if fmt is None:
        raise ExtractionError(
            f"Nieobsługiwany format pliku '{file_name}'. "
            f"Dostępne: {', '.join(sorted(set(FORMATS.values())))}."
        )
    if fmt == "PDF":
        text = _extract_pdf(data)
    elif fmt == "DOCX":
        text = _extract_docx(data)
    elif fmt == "HTML":
        text = _extract_html(data)
    elif fmt == "RTF":
        text = _extract_rtf(data)
    else:
        text = _extract_plain(data)
    if not text:
        raise ExtractionError("Nie udało się wyekstrahować żadnej treści z pliku.")
    return fmt, text
