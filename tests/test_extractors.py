"""Testy rejestratora ekstraktorów - wszystkie formaty + rozszerzalność."""

import io

import pytest
from docx import Document as DocxDocument

from backend.adapters import extractors
from backend.adapters.extractors import extract_text, format_of, labels, register
from backend.core.exceptions import ExtractionError


def make_pdf(text: str) -> bytes:
    """Minimalny, poprawny PDF z warstwą tekstową (bez zewnętrznych zależności)."""
    objects: list[bytes | None] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET".encode()
    objects[3] = (
        b"<< /Length " + str(len(stream)).encode() + b" >> stream\n" + stream + b"\nendstream"
    )

    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF"
    ).encode()
    return bytes(out)


def make_docx(paragraph: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Komórka A"
    table.rows[0].cells[1].text = "Komórka B"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------- happy path ----


@pytest.mark.parametrize(
    ("name", "data", "expected_label", "expected_fragment"),
    [
        ("dok.txt", "zwykły tekst".encode(), "TXT", "zwykły tekst"),
        ("dok.md", "# Nagłówek\n\ntreść".encode(), "Markdown", "Nagłówek"),
        ("dok.markdown", b"abc", "Markdown", "abc"),
        (
            "dok.html",
            "<html><script>x()</script><body><h1>Tytuł</h1></body></html>".encode(),
            "HTML",
            "Tytuł",
        ),
        ("dok.htm", "<p>treść strony</p>".encode(), "HTML", "treść strony"),
        (r"dok.rtf", rb"{\rtf1\ansi Hello RTF}", "RTF", "Hello RTF"),
        ("dok.pdf", make_pdf("Hello PDF"), "PDF", "Hello PDF"),
        ("dok.docx", make_docx("Akapit testowy"), "DOCX", "Akapit testowy"),
    ],
)
def test_extract_supported_formats(name, data, expected_label, expected_fragment):
    label, text = extract_text(name, data)
    assert label == expected_label
    assert expected_fragment in text


def test_docx_tables_are_extracted():
    _, text = extract_text("t.docx", make_docx("head"))
    assert "Komórka A | Komórka B" in text


def test_html_strips_script_and_style():
    html = b"<style>b{color:red}</style><script>alert(1)</script><p>widoczne</p>"
    _, text = extract_text("x.html", html)
    assert "widoczne" in text
    assert "color:red" not in text
    assert "alert" not in text


def test_text_is_cleaned():
    _, text = extract_text("x.txt", b"linia   \n\n\n\nkoniec\x00")
    assert "\x00" not in text
    assert "\n\n\n" not in text


# ------------------------------------------------------------------ błędy ----


def test_unsupported_format_rejected():
    with pytest.raises(ExtractionError) as exc:
        extract_text("program.exe", b"MZ")
    assert "Nieobsługiwany format" in str(exc.value)
    assert exc.value.status_code == 422


def test_no_extension_rejected():
    with pytest.raises(ExtractionError):
        extract_text("bezrozzerzenia", b"tekst")


def test_empty_txt_rejected():
    with pytest.raises(ExtractionError, match="pusty"):
        extract_text("pusty.txt", b"   \n  ")


def test_pdf_without_text_layer_rejected():
    pdf = make_pdf("")
    with pytest.raises(ExtractionError, match="warstwy tekstowej"):
        extract_text("skan.pdf", pdf)


def test_corrupt_pdf_rejected():
    with pytest.raises(ExtractionError, match="PDF"):
        extract_text("zly.pdf", b"to nie jest pdf")


def test_corrupt_docx_rejected():
    with pytest.raises(ExtractionError, match="DOCX"):
        extract_text("zly.docx", b"PK..uszkodzone")


# -------------------------------------------------------------- format_of ----


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("a.pdf", "PDF"),
        ("A.PDF", "PDF"),
        ("sciezka/do/pliku.MD", "Markdown"),
        ("archiwum.tar.gz", None),
        ("bezrozszerzenia", None),
        ("koniec.", None),
    ],
)
def test_format_of(name, expected):
    assert format_of(name) == expected


def test_supported_formats_and_accept_header():
    assert set(extractors.FORMATS) >= {".pdf", ".docx", ".txt", ".md", ".html", ".rtf"}
    assert "*.pdf" in extractors.ACCEPT
    assert extractors.ACCEPT.count(",") == len(extractors.FORMATS) - 1
    assert labels() == sorted(set(extractors.FORMATS.values()))


# ------------------------------------------------------------ rozszerzalność -


def test_register_new_format():
    snapshot = dict(extractors.REGISTRY)
    try:
        register(".csv", "CSV", lambda data: data.decode("utf-8"))
        assert format_of("dane.csv") == "CSV"
        assert extract_text("dane.csv", b"a,b,c") == ("CSV", "a,b,c")
        assert extractors.FORMATS[".csv"] == "CSV"
        assert "*.csv" in extractors.ACCEPT
    finally:
        extractors.REGISTRY.clear()
        extractors.REGISTRY.update(snapshot)
        extractors.FORMATS.clear()
        extractors.FORMATS.update({ext: lab for ext, (lab, _) in extractors.REGISTRY.items()})
        extractors.ACCEPT = ",".join(f"*{e}" for e in extractors.REGISTRY)


def test_register_requires_leading_dot():
    with pytest.raises(ValueError, match="kropki"):
        register("csv", "CSV", lambda data: data.decode())
